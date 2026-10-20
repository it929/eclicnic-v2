from rest_framework import status, permissions
from rest_framework.response import Response
from rest_framework.views import APIView
from decimal import Decimal
from django.db.models import Q

from queue_operations.models import RegFee
from patients.models import PatientPlan, Sponsor, PatientProfile
from IPD.models import Ward
from radio_lab.models import RadioLabInventory
from inventory.models import PharmacyTariff
from myAdmins.models import OtherService2, Packages, PackagesData
from users.models import User


# ==========================================================
# 1. SERVICE TARIFFS
# ==========================================================

# 1.1 Registration Fees
class RegistrationFeeView(APIView):
    permission_classes = [permissions.AllowAny]

    def get(self, request):
        plans = RegFee.objects.all().order_by('plan_name')
        data = [{'id': p.id, 'plan_name': p.plan_name, 'price': p.price} for p in plans]
        return Response({'total': len(data), 'plans': data})

    def post(self, request):
        plan_name = request.data.get('plan_name', '').strip()
        price = request.data.get('price', 0)

        if not plan_name:
            return Response({'detail': 'Plan name is required.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            price = int(price)
            if price < 0:
                return Response({'detail': 'Price cannot be negative.'}, status=status.HTTP_400_BAD_REQUEST)
        except (ValueError, TypeError):
            return Response({'detail': 'Invalid price format.'}, status=status.HTTP_400_BAD_REQUEST)

        if RegFee.objects.filter(plan_name__iexact=plan_name).exists():
            return Response({'detail': f'Plan "{plan_name}" already exists.'}, status=status.HTTP_400_BAD_REQUEST)

        obj = RegFee.objects.create(plan_name=plan_name, price=price)
        return Response({'detail': f'Plan "{obj.plan_name}" created successfully.', 'id': obj.id, 'plan_name': obj.plan_name, 'price': obj.price})

    def put(self, request, pk=None):
        try:
            plan = RegFee.objects.get(id=pk)
        except RegFee.DoesNotExist:
            return Response({'detail': 'Plan not found.'}, status=status.HTTP_404_NOT_FOUND)

        plan_name = request.data.get('plan_name', plan.plan_name).strip()
        price = request.data.get('price', plan.price)

        if RegFee.objects.filter(plan_name__iexact=plan_name).exclude(id=pk).exists():
            return Response({'detail': f'Another plan with name "{plan_name}" already exists.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            plan.price = int(price)
        except (ValueError, TypeError):
            return Response({'detail': 'Invalid price format.'}, status=status.HTTP_400_BAD_REQUEST)

        plan.plan_name = plan_name
        plan.save()
        return Response({'detail': 'Plan updated successfully.', 'id': plan.id, 'plan_name': plan.plan_name, 'price': plan.price})

    def delete(self, request, pk=None):
        try:
            plan = RegFee.objects.get(id=pk)
            plan.delete()
            return Response({'detail': 'Plan deleted successfully.'})
        except RegFee.DoesNotExist:
            return Response({'detail': 'Plan not found.'}, status=status.HTTP_404_NOT_FOUND)


# 1.2 Laboratory Charges
class LabChargesView(APIView):
    permission_classes = [permissions.AllowAny]

    def get(self, request):
        qs = RadioLabInventory.objects.filter(type='L').select_related('sponsor', 'plan', 'plan__sponsor').order_by('item')

        tariff_type = request.query_params.get('tariff_type', '').strip().lower()
        patient_id = request.query_params.get('patient')
        sponsor_id = request.query_params.get('sponsor')
        plan_id = request.query_params.get('plan')
        search = request.query_params.get('search', '').strip()

        # If explicitly requesting general out-of-pocket tariffs
        if tariff_type == 'general' or (sponsor_id and str(sponsor_id).lower() in ('general', 'self', 'cash')):
            qs = qs.filter(sponsor__isnull=True, plan__isnull=True)
        # If filtering for a specific patient
        elif patient_id or tariff_type == 'covered':
            if patient_id:
                try:
                    patient = PatientProfile.objects.select_related('sponsor', 'plan', 'category', 'sponsor__category').get(id=patient_id)
                    patient_sponsor = patient.sponsor
                    patient_plan = patient.plan

                    # Check if patient or sponsor has a Private category
                    is_private_patient = (
                        (patient.category and 'private' in str(patient.category.category).lower()) or
                        (patient_sponsor and patient_sponsor.category and 'private' in str(patient_sponsor.category.category).lower()) or
                        (patient_sponsor and ('private' in patient_sponsor.name.lower() or 'self' in patient_sponsor.name.lower())) or
                        (not patient_sponsor and not patient_plan)
                    )

                    if is_private_patient:
                        # Private sponsor maintains the same price as general!
                        # General price is always available for them because they pay out of pocket
                        qs = qs.filter(
                            (Q(sponsor__isnull=True) & Q(plan__isnull=True)) |
                            (Q(sponsor=patient_sponsor) if patient_sponsor else Q(pk__in=[]))
                        )
                    else:
                        # HMO Sponsor: uses covered plan tariffs (claims requiring HMO approval)
                        if not patient_sponsor and sponsor_id and str(sponsor_id).lower() not in ('all', 'general'):
                            try:
                                patient_sponsor = Sponsor.objects.get(id=sponsor_id)
                            except Sponsor.DoesNotExist:
                                pass
                        if not patient_plan and plan_id and str(plan_id).lower() not in ('all', 'general'):
                            try:
                                patient_plan = PatientPlan.objects.get(id=plan_id)
                            except PatientPlan.DoesNotExist:
                                pass

                        if tariff_type == 'general':
                            qs = qs.filter(sponsor__isnull=True, plan__isnull=True)
                        elif patient_sponsor and patient_plan:
                            qs = qs.filter(
                                (Q(sponsor=patient_sponsor) & Q(plan=patient_plan)) |
                                Q(plan=patient_plan) |
                                (Q(sponsor=patient_sponsor) & Q(plan__isnull=True))
                            )
                        elif patient_plan:
                            qs = qs.filter(plan=patient_plan)
                        elif patient_sponsor:
                            qs = qs.filter(Q(sponsor=patient_sponsor) | Q(plan__sponsor=patient_sponsor))
                        else:
                            qs = qs.filter(sponsor__isnull=True, plan__isnull=True)
                except PatientProfile.DoesNotExist:
                    if sponsor_id and str(sponsor_id).lower() != 'all':
                        qs = qs.filter(Q(sponsor_id=sponsor_id) | Q(plan__sponsor_id=sponsor_id))
                    if plan_id and str(plan_id).lower() != 'all':
                        qs = qs.filter(plan_id=plan_id)
            else:
                if sponsor_id and str(sponsor_id).lower() != 'all':
                    qs = qs.filter(Q(sponsor_id=sponsor_id) | Q(plan__sponsor_id=sponsor_id))
                if plan_id and str(plan_id).lower() != 'all':
                    qs = qs.filter(plan_id=plan_id)
        else:
            if sponsor_id and str(sponsor_id).lower() != 'all':
                qs = qs.filter(Q(sponsor_id=sponsor_id) | Q(plan__sponsor_id=sponsor_id))
            if plan_id and str(plan_id).lower() != 'all':
                qs = qs.filter(plan_id=plan_id)

        if search:
            qs = qs.filter(
                Q(item__icontains=search) |
                Q(item_id__icontains=search)
            )

        data = []
        for item in qs:
            eff_sponsor = item.sponsor or (item.plan.sponsor if item.plan else None)
            is_gen = bool(not eff_sponsor and not item.plan)
            data.append({
                'id': item.id,
                'item': item.item,
                'item_id': item.item_id or '',
                'rate': float(item.rate),
                'sponsor_id': eff_sponsor.id if eff_sponsor else None,
                'sponsor_name': eff_sponsor.name if eff_sponsor else 'General / Self-Pay',
                'sponsor_code': eff_sponsor.code if eff_sponsor else 'GENERAL',
                'plan_id': item.plan.id if item.plan else None,
                'plan_name': item.plan.plan if item.plan else ('General Price' if is_gen else 'All Plans'),
                'plan_code': item.plan.code if item.plan else ('GENERAL' if is_gen else ''),
                'is_general': is_gen,
                'payment_label': 'Pay from your purse' if is_gen else 'Covered by Plan',
            })
        return Response({'total': len(data), 'charges': data, 'results': data})

    def post(self, request):
        item_name = request.data.get('item', '').strip() or request.data.get('item_name', '').strip()
        item_id = request.data.get('item_id', '').strip()
        rate = request.data.get('rate', 0)
        sponsor_id = request.data.get('sponsor_id') or request.data.get('sponsor')
        plan_id = request.data.get('plan_id') or request.data.get('plan')

        if not item_name:
            return Response({'detail': 'Test item name required.'}, status=status.HTTP_400_BAD_REQUEST)

        is_general = (
            not sponsor_id or
            str(sponsor_id).lower() in ('', 'general', 'self', 'cash', 'none', 'null')
        )

        plan = None
        sponsor = None

        if not is_general:
            if plan_id and str(plan_id).lower() not in ('all', 'general', 'none'):
                try:
                    plan = PatientPlan.objects.select_related('sponsor').get(id=plan_id)
                except PatientPlan.DoesNotExist:
                    return Response({'detail': 'Selected Plan not found.'}, status=status.HTTP_404_NOT_FOUND)

            if sponsor_id and str(sponsor_id).lower() not in ('all', 'general', 'none'):
                try:
                    sponsor = Sponsor.objects.get(id=sponsor_id)
                except Sponsor.DoesNotExist:
                    return Response({'detail': 'Selected Sponsor not found.'}, status=status.HTTP_404_NOT_FOUND)
            elif plan and plan.sponsor:
                sponsor = plan.sponsor

        staff = request.user if request.user.is_authenticated else User.objects.first()

        obj = RadioLabInventory.objects.create(
            item=item_name,
            item_id=item_id or f"LAB-{RadioLabInventory.objects.count()+1:03d}",
            rate=Decimal(str(rate)),
            type='L',
            sponsor=sponsor,
            plan=plan,
            staff=staff
        )

        eff_sponsor = obj.sponsor or (obj.plan.sponsor if obj.plan else None)
        is_gen = bool(not eff_sponsor and not obj.plan)
        return Response({
            'detail': f'Lab charge "{obj.item}" created successfully.',
            'id': obj.id,
            'item': obj.item,
            'item_id': obj.item_id,
            'rate': float(obj.rate),
            'sponsor_id': eff_sponsor.id if eff_sponsor else None,
            'sponsor_name': eff_sponsor.name if eff_sponsor else 'General / Self-Pay',
            'sponsor_code': eff_sponsor.code if eff_sponsor else 'GENERAL',
            'plan_id': obj.plan.id if obj.plan else None,
            'plan_name': obj.plan.plan if obj.plan else ('General Price' if is_gen else 'All Plans'),
            'plan_code': obj.plan.code if obj.plan else ('GENERAL' if is_gen else ''),
            'is_general': is_gen,
            'payment_label': 'Pay from your purse' if is_gen else 'Covered by Plan',
        }, status=status.HTTP_201_CREATED)

    def put(self, request, pk=None):
        try:
            obj = RadioLabInventory.objects.select_related('sponsor', 'plan').get(id=pk, type='L')
        except RadioLabInventory.DoesNotExist:
            return Response({'detail': 'Lab charge not found.'}, status=status.HTTP_404_NOT_FOUND)

        if 'item' in request.data:
            obj.item = request.data['item'].strip()
        if 'rate' in request.data:
            obj.rate = Decimal(str(request.data['rate']))
        if 'item_id' in request.data:
            obj.item_id = request.data['item_id'].strip()

        sponsor_id = request.data.get('sponsor_id') or request.data.get('sponsor')
        if sponsor_id is not None:
            if sponsor_id and str(sponsor_id).lower() != 'all':
                try:
                    obj.sponsor = Sponsor.objects.get(id=sponsor_id)
                except Sponsor.DoesNotExist:
                    pass
            else:
                obj.sponsor = None

        plan_id = request.data.get('plan_id') or request.data.get('plan')
        if plan_id is not None:
            if plan_id and str(plan_id).lower() != 'all':
                try:
                    obj.plan = PatientPlan.objects.get(id=plan_id)
                    if obj.plan.sponsor and not obj.sponsor:
                        obj.sponsor = obj.plan.sponsor
                except PatientPlan.DoesNotExist:
                    pass
            else:
                obj.plan = None

        obj.save()
        eff_sponsor = obj.sponsor or (obj.plan.sponsor if obj.plan else None)
        return Response({
            'detail': 'Lab charge updated successfully.',
            'id': obj.id,
            'item': obj.item,
            'item_id': obj.item_id,
            'rate': float(obj.rate),
            'sponsor_id': eff_sponsor.id if eff_sponsor else None,
            'sponsor_name': eff_sponsor.name if eff_sponsor else 'General',
            'plan_id': obj.plan.id if obj.plan else None,
            'plan_name': obj.plan.plan if obj.plan else 'All Plans',
        })

    def delete(self, request, pk=None):
        try:
            RadioLabInventory.objects.filter(id=pk, type='L').delete()
            return Response({'detail': 'Lab charge deleted.'})
        except Exception as e:
            return Response({'detail': str(e)}, status=status.HTTP_400_BAD_REQUEST)


# 1.2.1 Download Laboratory Tariff Template
class DownloadLabTariffTemplateView(APIView):
    permission_classes = [permissions.AllowAny]

    def get(self, request):
        from django.http import HttpResponse
        import io, csv

        sponsor_id = request.query_params.get('sponsor')
        plan_id = request.query_params.get('plan')

        is_general = str(sponsor_id).strip().lower() in ('general', 'self', 'cash', 'none')

        if is_general:
            sponsor_code = "GENERAL"
            plan_code = "GENERAL"
            sponsor_name = "General / Out-of-Pocket"
            plan_name = "General Price (Pay from your purse)"
        else:
            sponsor_code = "AMH"
            plan_code = "AIICO SLIVER"
            sponsor_name = "AIICO MULTISHIELD (AMH)"
            plan_name = "AIICO SLIVER"

            if sponsor_id:
                try:
                    s = Sponsor.objects.get(id=sponsor_id)
                    sponsor_code = s.code or s.name
                    sponsor_name = s.name
                except Sponsor.DoesNotExist:
                    pass

            if plan_id:
                try:
                    p = PatientPlan.objects.select_related('sponsor').get(id=plan_id)
                    plan_code = p.code or p.plan
                    plan_name = p.plan
                    if p.sponsor:
                        sponsor_code = p.sponsor.code or p.sponsor.name
                        sponsor_name = p.sponsor.name
                except PatientPlan.DoesNotExist:
                    pass

        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(['item_name', 'test_code', 'rate', 'sponsor_code', 'plan_code', 'notes'])

        notes_text = "General Tariff - Pay from your purse" if is_general else f"Covered under {plan_name}"

        samples = [
            ['Full Blood Count (FBC / CBC)', 'LAB-001', '3500.00', sponsor_code, plan_code, notes_text],
            ['Malaria Parasite (MP by Microscopy)', 'LAB-002', '1500.00', sponsor_code, plan_code, notes_text],
            ['Widal Agglutination Test (Typhoid)', 'LAB-003', '2000.00', sponsor_code, plan_code, notes_text],
            ['Urinalysis (10-Parameter Strip)', 'LAB-004', '1200.00', sponsor_code, plan_code, notes_text],
            ['Fasting Blood Sugar (FBS)', 'LAB-005', '1000.00', sponsor_code, plan_code, notes_text],
            ['Lipid Profile (Full)', 'LAB-006', '6500.00', sponsor_code, plan_code, notes_text],
            ['Liver Function Test (LFT)', 'LAB-007', '5500.00', sponsor_code, plan_code, notes_text],
            ['Kidney Function Test (Electrolytes, Urea, Creatinine)', 'LAB-008', '6000.00', sponsor_code, plan_code, notes_text],
            ['Stool Microscopy, Culture & Sensitivity', 'LAB-009', '4000.00', sponsor_code, plan_code, notes_text],
            ['Hepatitis B Surface Antigen (HBsAg Screening)', 'LAB-010', '2500.00', sponsor_code, plan_code, notes_text],
        ]
        for row in samples:
            writer.writerow(row)

        csv_content = output.getvalue()
        response = HttpResponse(csv_content, content_type='text/csv; charset=utf-8')
        safe_sponsor = str(sponsor_code).replace(' ', '_').replace('/', '_')
        safe_plan = str(plan_code).replace(' ', '_').replace('/', '_')
        filename = "lab_tariffs_template_general.csv" if is_general else f"lab_tariffs_template_{safe_sponsor}_{safe_plan}.csv"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


# 1.2.2 Bulk Upload Laboratory Tariffs
class BulkUploadLabTariffView(APIView):
    permission_classes = [permissions.AllowAny]

    def post(self, request):
        import io, csv
        default_sponsor_id = request.data.get('sponsor_id') or request.data.get('default_sponsor_id')
        default_plan_id = request.data.get('plan_id') or request.data.get('default_plan_id')

        is_general_default = str(default_sponsor_id).strip().lower() in ('general', 'self', 'cash', 'none')

        default_sponsor = None
        if not is_general_default and default_sponsor_id and str(default_sponsor_id).lower() not in ('', 'all', 'null', 'none'):
            try:
                default_sponsor = Sponsor.objects.get(id=default_sponsor_id)
            except Sponsor.DoesNotExist:
                pass

        default_plan = None
        if not is_general_default and default_plan_id and str(default_plan_id).lower() not in ('', 'all', 'null', 'none'):
            try:
                default_plan = PatientPlan.objects.select_related('sponsor').get(id=default_plan_id)
                if not default_sponsor and default_plan.sponsor:
                    default_sponsor = default_plan.sponsor
            except PatientPlan.DoesNotExist:
                pass

        raw_rows = []
        uploaded_file = request.FILES.get('file')

        if uploaded_file:
            filename = uploaded_file.name.lower()
            try:
                if filename.endswith('.csv'):
                    decoded_file = uploaded_file.read().decode('utf-8-sig', errors='ignore')
                    reader = csv.DictReader(io.StringIO(decoded_file))
                    for row in reader:
                        raw_rows.append(row)
                elif filename.endswith(('.xlsx', '.xls')):
                    import pandas as pd
                    df = pd.read_excel(io.BytesIO(uploaded_file.read()))
                    df = df.where(pd.notnull(df), None)
                    raw_rows = df.to_dict(orient='records')
                else:
                    return Response(
                        {'detail': 'Unsupported file format. Please upload a CSV (.csv) or Excel (.xlsx, .xls) file.'},
                        status=status.HTTP_400_BAD_REQUEST
                    )
            except Exception as e:
                return Response(
                    {'detail': f'Error reading uploaded file: {str(e)}'},
                    status=status.HTTP_400_BAD_REQUEST
                )
        elif 'items' in request.data:
            items = request.data.get('items', [])
            if isinstance(items, list):
                raw_rows = items
        else:
            return Response(
                {'detail': 'No file or items array provided for bulk upload.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        if not raw_rows:
            return Response(
                {'detail': 'The uploaded file contains no data rows.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        all_sponsors = list(Sponsor.objects.all())
        all_plans = list(PatientPlan.objects.select_related('sponsor').all())

        sponsor_by_code = {s.code.strip().upper(): s for s in all_sponsors if s.code}
        sponsor_by_name = {s.name.strip().upper(): s for s in all_sponsors if s.name}
        sponsor_by_id = {str(s.id): s for s in all_sponsors}

        plan_by_code = {p.code.strip().upper(): p for p in all_plans if p.code}
        plan_by_name = {p.plan.strip().upper(): p for p in all_plans if p.plan}
        plan_by_id = {str(p.id): p for p in all_plans}

        staff = request.user if request.user.is_authenticated else User.objects.first()

        created_count = 0
        updated_count = 0
        skipped = []

        counter = RadioLabInventory.objects.count()

        for idx, row in enumerate(raw_rows, start=1):
            if not isinstance(row, dict):
                continue

            clean_row = {str(k).strip().lower(): v for k, v in row.items() if k is not None}

            item_name = (
                clean_row.get('item_name')
                or clean_row.get('test_name')
                or clean_row.get('laboratory_test')
                or clean_row.get('lab_test')
                or clean_row.get('test')
                or clean_row.get('item')
                or clean_row.get('name')
            )
            if not item_name or not str(item_name).strip():
                skipped.append(f"Row {idx}: Missing test item name")
                continue

            item_name = str(item_name).strip()

            rate_raw = (
                clean_row.get('rate')
                or clean_row.get('tariff_rate')
                or clean_row.get('price')
                or clean_row.get('amount')
                or clean_row.get('fee')
            )
            try:
                rate_val = Decimal(str(rate_raw).replace(',', '').replace('₦', '').strip())
                if rate_val < 0:
                    rate_val = Decimal('0.00')
            except Exception:
                skipped.append(f"Row {idx} ({item_name}): Invalid rate '{rate_raw}'")
                continue

            test_code = (
                clean_row.get('test_code')
                or clean_row.get('item_id')
                or clean_row.get('code')
                or clean_row.get('lab_code')
            )
            if not test_code or not str(test_code).strip():
                counter += 1
                test_code = f"LAB-{counter:03d}"
            else:
                test_code = str(test_code).strip()

            # Resolve Sponsor
            row_sponsor_val = clean_row.get('sponsor_code') or clean_row.get('sponsor') or clean_row.get('sponsor_name')
            row_sponsor = None
            row_is_general = False

            if row_sponsor_val and str(row_sponsor_val).strip():
                raw_s = str(row_sponsor_val).strip()
                if raw_s.upper() in ('GENERAL', 'SELF', 'SELF-PAY', 'CASH', 'OUT-OF-POCKET', 'NO HMO', 'GENERAL TARIFF', 'NONE'):
                    row_is_general = True
                else:
                    row_sponsor = (
                        sponsor_by_code.get(raw_s.upper())
                        or sponsor_by_name.get(raw_s.upper())
                        or sponsor_by_id.get(raw_s)
                    )

            target_sponsor = None if (row_is_general or is_general_default) else (row_sponsor or default_sponsor)

            # Resolve Plan
            row_plan_val = clean_row.get('plan_code') or clean_row.get('plan') or clean_row.get('plan_name')
            row_plan = None
            if not row_is_general and not is_general_default and row_plan_val and str(row_plan_val).strip():
                raw_p = str(row_plan_val).strip()
                if raw_p.upper() not in ('GENERAL', 'ALL', 'NONE', 'NO HMO', 'PAY FROM YOUR PURSE', 'PURSE'):
                    row_plan = (
                        plan_by_code.get(raw_p.upper())
                        or plan_by_name.get(raw_p.upper())
                        or plan_by_id.get(raw_p)
                    )

            target_plan = None if (row_is_general or is_general_default) else (row_plan or default_plan)
            if target_plan and target_plan.sponsor and not target_sponsor:
                target_sponsor = target_plan.sponsor

            # Check if this row is a General Tariff (Self-Pay / Pay from your purse)
            is_general_row = is_general_default or row_is_general or (
                target_sponsor is None and target_plan is None and (
                    default_sponsor_id in ('general', 'self', 'cash', 'none', '', None) or
                    clean_row.get('sponsor_code', '').strip().upper() in ('GENERAL', 'SELF', 'CASH', '')
                )
            )

            if not is_general_row and not target_sponsor and not target_plan:
                skipped.append(f"Row {idx} ({item_name}): Could not identify Sponsor or Plan. Select an HMO or choose General Tariff.")
                continue

            if is_general_row:
                # General Tariff: sponsor=None, plan=None
                existing = RadioLabInventory.objects.filter(
                    type='L',
                    item__iexact=item_name,
                    sponsor__isnull=True,
                    plan__isnull=True
                ).first()

                if not existing and test_code:
                    existing = RadioLabInventory.objects.filter(
                        type='L',
                        item_id__iexact=test_code,
                        sponsor__isnull=True,
                        plan__isnull=True
                    ).first()

                if existing:
                    existing.rate = rate_val
                    existing.item = item_name
                    existing.item_id = test_code
                    existing.sponsor = None
                    existing.plan = None
                    existing.staff = staff
                    existing.save()
                    updated_count += 1
                else:
                    RadioLabInventory.objects.create(
                        item=item_name,
                        item_id=test_code,
                        rate=rate_val,
                        type='L',
                        sponsor=None,
                        plan=None,
                        staff=staff
                    )
                    created_count += 1
            else:
                # HMO Tariff: linked to specific Sponsor and/or Plan
                existing = RadioLabInventory.objects.filter(
                    type='L',
                    item__iexact=item_name,
                    sponsor=target_sponsor,
                    plan=target_plan
                ).first()

                if not existing and test_code:
                    existing = RadioLabInventory.objects.filter(
                        type='L',
                        item_id__iexact=test_code,
                        sponsor=target_sponsor,
                        plan=target_plan
                    ).first()

                if existing:
                    existing.rate = rate_val
                    existing.item = item_name
                    existing.item_id = test_code
                    existing.staff = staff
                    existing.save()
                    updated_count += 1
                else:
                    RadioLabInventory.objects.create(
                        item=item_name,
                        item_id=test_code,
                        rate=rate_val,
                        type='L',
                        sponsor=target_sponsor,
                        plan=target_plan,
                        staff=staff
                    )
                    created_count += 1

        msg = f"Bulk upload completed: {created_count} laboratory tariffs created, {updated_count} updated."
        if is_general_default:
            msg = f"Bulk upload for General Laboratory Tariffs completed: {created_count} created, {updated_count} updated."
        if skipped:
            msg += f" {len(skipped)} rows were skipped."

        return Response({
            'detail': msg,
            'created': created_count,
            'updated': updated_count,
            'skipped': skipped,
            'total_processed': created_count + updated_count + len(skipped),
        }, status=status.HTTP_200_OK)


# 1.3 Radiology Charges
class RadiologyChargesView(APIView):
    permission_classes = [permissions.AllowAny]

    def get(self, request):
        qs = RadioLabInventory.objects.filter(type='R').select_related('sponsor', 'plan', 'plan__sponsor').order_by('item')

        patient_id = request.query_params.get('patient')
        sponsor_id = request.query_params.get('sponsor')
        plan_id = request.query_params.get('plan')
        search = request.query_params.get('search', '').strip()

        # If filtering for a specific patient (e.g. from Doctor Consultation)
        if patient_id:
            try:
                patient = PatientProfile.objects.select_related('sponsor', 'plan').get(id=patient_id)
                patient_sponsor = patient.sponsor
                patient_plan = patient.plan

                if patient_sponsor and patient_plan:
                    patient_filter = (
                        (Q(sponsor=patient_sponsor) & Q(plan=patient_plan)) |
                        Q(plan=patient_plan) |
                        (Q(sponsor=patient_sponsor) & Q(plan__isnull=True)) |
                        (Q(sponsor__isnull=True) & Q(plan__isnull=True))
                    )
                elif patient_plan:
                    patient_filter = (
                        Q(plan=patient_plan) |
                        (Q(sponsor__isnull=True) & Q(plan__isnull=True))
                    )
                elif patient_sponsor:
                    patient_filter = (
                        Q(sponsor=patient_sponsor) |
                        Q(plan__sponsor=patient_sponsor) |
                        (Q(sponsor__isnull=True) & Q(plan__isnull=True))
                    )
                else:
                    patient_filter = (Q(sponsor__isnull=True) & Q(plan__isnull=True))

                qs = qs.filter(patient_filter)
            except PatientProfile.DoesNotExist:
                pass
        else:
            if sponsor_id and str(sponsor_id).lower() != 'all':
                qs = qs.filter(Q(sponsor_id=sponsor_id) | Q(plan__sponsor_id=sponsor_id))
            if plan_id and str(plan_id).lower() != 'all':
                qs = qs.filter(plan_id=plan_id)

        if search:
            qs = qs.filter(
                Q(item__icontains=search) |
                Q(item_id__icontains=search)
            )

        data = []
        for item in qs:
            eff_sponsor = item.sponsor or (item.plan.sponsor if item.plan else None)
            data.append({
                'id': item.id,
                'item': item.item,
                'item_id': item.item_id or '',
                'rate': float(item.rate),
                'sponsor_id': eff_sponsor.id if eff_sponsor else None,
                'sponsor_name': eff_sponsor.name if eff_sponsor else 'General / Self',
                'sponsor_code': eff_sponsor.code if eff_sponsor else '',
                'plan_id': item.plan.id if item.plan else None,
                'plan_name': item.plan.plan if item.plan else 'All Plans',
            })
        return Response({'total': len(data), 'charges': data})

    def post(self, request):
        item_name = request.data.get('item', '').strip()
        item_id = request.data.get('item_id', '').strip()
        rate = request.data.get('rate', 0)
        sponsor_id = request.data.get('sponsor_id') or request.data.get('sponsor')
        plan_id = request.data.get('plan_id') or request.data.get('plan')

        if not item_name:
            return Response({'detail': 'Scan item name required.'}, status=status.HTTP_400_BAD_REQUEST)

        sponsor = None
        if sponsor_id and str(sponsor_id).lower() != 'all':
            try:
                sponsor = Sponsor.objects.get(id=sponsor_id)
            except Sponsor.DoesNotExist:
                return Response({'detail': 'Sponsor not found.'}, status=status.HTTP_404_NOT_FOUND)

        plan = None
        if plan_id and str(plan_id).lower() != 'all':
            try:
                plan = PatientPlan.objects.get(id=plan_id)
                if plan.sponsor and not sponsor:
                    sponsor = plan.sponsor
            except PatientPlan.DoesNotExist:
                return Response({'detail': 'Plan not found.'}, status=status.HTTP_404_NOT_FOUND)

        staff = request.user if request.user.is_authenticated else User.objects.first()

        try:
            rate_decimal = Decimal(str(rate))
            if rate_decimal < 0:
                rate_decimal = Decimal('0.00')
        except Exception:
            return Response({'detail': 'Invalid rate value.'}, status=status.HTTP_400_BAD_REQUEST)

        obj = RadioLabInventory.objects.create(
            item=item_name,
            item_id=item_id or f"RAD-{RadioLabInventory.objects.count()+1:03d}",
            rate=rate_decimal,
            type='R',
            sponsor=sponsor,
            plan=plan,
            staff=staff
        )
        eff_sponsor = obj.sponsor or (obj.plan.sponsor if obj.plan else None)
        return Response({
            'detail': f'Radiology charge "{obj.item}" created successfully.',
            'id': obj.id,
            'item': obj.item,
            'item_id': obj.item_id,
            'rate': float(obj.rate),
            'sponsor_id': eff_sponsor.id if eff_sponsor else None,
            'sponsor_name': eff_sponsor.name if eff_sponsor else 'General',
            'plan_id': obj.plan.id if obj.plan else None,
            'plan_name': obj.plan.plan if obj.plan else 'All Plans',
        }, status=status.HTTP_201_CREATED)

    def put(self, request, pk=None):
        try:
            obj = RadioLabInventory.objects.get(id=pk, type='R')
        except RadioLabInventory.DoesNotExist:
            return Response({'detail': 'Radiology charge not found.'}, status=status.HTTP_404_NOT_FOUND)

        if 'item' in request.data:
            obj.item = request.data['item'].strip()
        if 'rate' in request.data:
            try:
                obj.rate = Decimal(str(request.data['rate']))
            except Exception:
                return Response({'detail': 'Invalid rate format.'}, status=status.HTTP_400_BAD_REQUEST)
        if 'item_id' in request.data:
            obj.item_id = request.data['item_id'].strip()

        sponsor_id = request.data.get('sponsor_id') or request.data.get('sponsor')
        if sponsor_id is not None:
            if sponsor_id and str(sponsor_id).lower() != 'all':
                try:
                    obj.sponsor = Sponsor.objects.get(id=sponsor_id)
                except Sponsor.DoesNotExist:
                    return Response({'detail': 'Sponsor not found.'}, status=status.HTTP_404_NOT_FOUND)
            else:
                obj.sponsor = None

        plan_id = request.data.get('plan_id') or request.data.get('plan')
        if plan_id is not None:
            if plan_id and str(plan_id).lower() != 'all':
                try:
                    obj.plan = PatientPlan.objects.get(id=plan_id)
                    if obj.plan.sponsor and not obj.sponsor:
                        obj.sponsor = obj.plan.sponsor
                except PatientPlan.DoesNotExist:
                    pass
            else:
                obj.plan = None

        obj.save()
        eff_sponsor = obj.sponsor or (obj.plan.sponsor if obj.plan else None)
        return Response({
            'detail': 'Radiology charge updated successfully.',
            'id': obj.id,
            'item': obj.item,
            'item_id': obj.item_id,
            'rate': float(obj.rate),
            'sponsor_id': eff_sponsor.id if eff_sponsor else None,
            'sponsor_name': eff_sponsor.name if eff_sponsor else 'General',
            'plan_id': obj.plan.id if obj.plan else None,
            'plan_name': obj.plan.plan if obj.plan else 'All Plans',
        })

    def delete(self, request, pk=None):
        try:
            RadioLabInventory.objects.filter(id=pk, type='R').delete()
            return Response({'detail': 'Radiology charge deleted.'})
        except Exception as e:
            return Response({'detail': str(e)}, status=status.HTTP_400_BAD_REQUEST)


# 1.3.1 Download Radiology Tariff Template
class DownloadRadiologyTariffTemplateView(APIView):
    permission_classes = [permissions.AllowAny]

    def get(self, request):
        from django.http import HttpResponse
        import io, csv

        sponsor_id = request.query_params.get('sponsor')
        plan_id = request.query_params.get('plan')

        sponsor_code = "AMH"
        plan_code = "AIICO SLIVER"
        sponsor_name = "AIICO MULTISHIELD (AMH)"
        plan_name = "AIICO SLIVER"

        if sponsor_id:
            try:
                s = Sponsor.objects.get(id=sponsor_id)
                sponsor_code = s.code or s.name
                sponsor_name = s.name
            except Sponsor.DoesNotExist:
                pass

        if plan_id:
            try:
                p = PatientPlan.objects.select_related('sponsor').get(id=plan_id)
                plan_code = p.code or p.plan
                plan_name = p.plan
                if p.sponsor:
                    sponsor_code = p.sponsor.code or p.sponsor.name
                    sponsor_name = p.sponsor.name
            except PatientPlan.DoesNotExist:
                pass

        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(['item_name', 'procedure_code', 'rate', 'sponsor_code', 'plan_code', 'notes'])

        samples = [
            ['Chest X-Ray (PA View)', 'RAD-001', '5000.00', sponsor_code, plan_code, f'Covered under {plan_name}'],
            ['Abdominal Ultrasound Scan', 'RAD-002', '6000.00', sponsor_code, plan_code, f'Covered under {plan_name}'],
            ['Pelvic Ultrasound Scan', 'RAD-003', '5500.00', sponsor_code, plan_code, f'Covered under {plan_name}'],
            ['Obstetric Ultrasound (Pregnancy Scan)', 'RAD-004', '4500.00', sponsor_code, plan_code, f'Covered under {plan_name}'],
            ['Lumbosacral Spine X-Ray (AP & Lateral)', 'RAD-005', '7500.00', sponsor_code, plan_code, f'Covered under {plan_name}'],
            ['Brain CT Scan (Non-Contrast)', 'RAD-006', '35000.00', sponsor_code, plan_code, f'Covered under {plan_name}'],
            ['Echocardiography (ECHO)', 'RAD-007', '20000.00', sponsor_code, plan_code, f'Covered under {plan_name}'],
            ['Electrocardiogram (ECG / EKG)', 'RAD-008', '4000.00', sponsor_code, plan_code, f'Covered under {plan_name}'],
            ['Mammography (Bilateral)', 'RAD-009', '15000.00', sponsor_code, plan_code, f'Covered under {plan_name}'],
            ['Knee Joint X-Ray (AP & Lateral)', 'RAD-010', '6000.00', sponsor_code, plan_code, f'Covered under {plan_name}'],
        ]
        for row in samples:
            writer.writerow(row)

        response = HttpResponse(output.getvalue(), content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="radiology_tariffs_template.csv"'
        return response


# 1.3.2 Bulk Upload Radiology Tariffs
class BulkUploadRadiologyTariffView(APIView):
    permission_classes = [permissions.AllowAny]

    def post(self, request):
        import csv, io
        uploaded_file = request.FILES.get('file')
        default_sponsor_id = request.data.get('sponsor_id') or request.data.get('sponsor')
        default_plan_id = request.data.get('plan_id') or request.data.get('plan')

        default_sponsor = None
        if default_sponsor_id and str(default_sponsor_id).lower() != 'all':
            try:
                default_sponsor = Sponsor.objects.get(id=default_sponsor_id)
            except Sponsor.DoesNotExist:
                pass

        default_plan = None
        if default_plan_id and str(default_plan_id).lower() != 'all':
            try:
                default_plan = PatientPlan.objects.get(id=default_plan_id)
                if default_plan.sponsor and not default_sponsor:
                    default_sponsor = default_plan.sponsor
            except PatientPlan.DoesNotExist:
                pass

        raw_rows = []

        if uploaded_file:
            fname = uploaded_file.name.lower()
            if fname.endswith('.csv'):
                try:
                    file_data = uploaded_file.read().decode('utf-8', errors='replace')
                    reader = csv.DictReader(io.StringIO(file_data))
                    for row in reader:
                        raw_rows.append(row)
                except Exception as e:
                    return Response({'detail': f'Could not parse CSV file: {str(e)}'}, status=status.HTTP_400_BAD_REQUEST)
            elif fname.endswith(('.xls', '.xlsx')):
                try:
                    import openpyxl
                    wb = openpyxl.load_workbook(uploaded_file, data_only=True)
                    ws = wb.active
                    headers = [str(cell.value or '').strip() for cell in ws[1]]
                    for row in ws.iter_rows(min_row=2, values_only=True):
                        if any(v is not None for v in row):
                            row_dict = {headers[i]: row[i] for i in range(len(headers)) if i < len(row)}
                            raw_rows.append(row_dict)
                except Exception as e:
                    return Response({'detail': f'Could not parse Excel file: {str(e)}'}, status=status.HTTP_400_BAD_REQUEST)
            else:
                return Response(
                    {'detail': 'Unsupported file format. Please upload a .csv or .xlsx file.'},
                    status=status.HTTP_400_BAD_REQUEST
                )
        elif 'items' in request.data:
            items = request.data.get('items', [])
            if isinstance(items, list):
                raw_rows = items
        else:
            return Response(
                {'detail': 'No file or items array provided for bulk upload.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        if not raw_rows:
            return Response(
                {'detail': 'The uploaded file contains no data rows.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        all_sponsors = list(Sponsor.objects.all())
        all_plans = list(PatientPlan.objects.select_related('sponsor').all())

        sponsor_by_code = {s.code.strip().upper(): s for s in all_sponsors if s.code}
        sponsor_by_name = {s.name.strip().upper(): s for s in all_sponsors if s.name}
        sponsor_by_id = {str(s.id): s for s in all_sponsors}

        plan_by_code = {p.code.strip().upper(): p for p in all_plans if p.code}
        plan_by_name = {p.plan.strip().upper(): p for p in all_plans if p.plan}
        plan_by_id = {str(p.id): p for p in all_plans}

        staff = request.user if request.user.is_authenticated else User.objects.first()

        created_count = 0
        updated_count = 0
        skipped = []

        counter = RadioLabInventory.objects.count()

        for idx, row in enumerate(raw_rows, start=1):
            if not isinstance(row, dict):
                continue

            clean_row = {str(k).strip().lower(): v for k, v in row.items() if k is not None}

            item_name = (
                clean_row.get('item_name')
                or clean_row.get('procedure_name')
                or clean_row.get('test_name')
                or clean_row.get('radiology_procedure')
                or clean_row.get('scan_name')
                or clean_row.get('scan')
                or clean_row.get('procedure')
                or clean_row.get('item')
                or clean_row.get('name')
            )
            if not item_name or not str(item_name).strip():
                skipped.append(f"Row {idx}: Missing scan procedure name")
                continue

            item_name = str(item_name).strip()

            rate_raw = (
                clean_row.get('rate')
                or clean_row.get('tariff_rate')
                or clean_row.get('price')
                or clean_row.get('amount')
                or clean_row.get('fee')
            )
            try:
                rate_val = Decimal(str(rate_raw).replace(',', '').replace('₦', '').strip())
                if rate_val < 0:
                    rate_val = Decimal('0.00')
            except Exception:
                skipped.append(f"Row {idx} ({item_name}): Invalid rate '{rate_raw}'")
                continue

            proc_code = (
                clean_row.get('procedure_code')
                or clean_row.get('test_code')
                or clean_row.get('item_id')
                or clean_row.get('code')
            )
            if proc_code:
                proc_code = str(proc_code).strip().upper()
            else:
                counter += 1
                proc_code = f"RAD-{counter:03d}"

            # Resolve Sponsor
            row_sponsor_val = clean_row.get('sponsor_code') or clean_row.get('sponsor') or clean_row.get('sponsor_name')
            row_sponsor = None
            if row_sponsor_val and str(row_sponsor_val).strip():
                raw_s = str(row_sponsor_val).strip()
                row_sponsor = (
                    sponsor_by_code.get(raw_s.upper())
                    or sponsor_by_name.get(raw_s.upper())
                    or sponsor_by_id.get(raw_s)
                )

            target_sponsor = row_sponsor or default_sponsor

            # Resolve Plan
            row_plan_val = clean_row.get('plan_code') or clean_row.get('plan') or clean_row.get('plan_name')
            row_plan = None
            if row_plan_val and str(row_plan_val).strip():
                raw_p = str(row_plan_val).strip()
                row_plan = (
                    plan_by_code.get(raw_p.upper())
                    or plan_by_name.get(raw_p.upper())
                    or plan_by_id.get(raw_p)
                )

            target_plan = row_plan or default_plan
            if target_plan and target_plan.sponsor and not target_sponsor:
                target_sponsor = target_plan.sponsor

            if not target_sponsor and not target_plan:
                skipped.append(f"Row {idx} ({item_name}): Could not identify Sponsor or Plan")
                continue

            # Check existing by item + sponsor + plan
            existing = RadioLabInventory.objects.filter(
                type='R',
                item__iexact=item_name,
                sponsor=target_sponsor,
                plan=target_plan
            ).first()

            if not existing and proc_code:
                existing = RadioLabInventory.objects.filter(
                    type='R',
                    item_id__iexact=proc_code,
                    sponsor=target_sponsor,
                    plan=target_plan
                ).first()

            if existing:
                existing.rate = rate_val
                existing.item = item_name
                existing.item_id = proc_code
                existing.staff = staff
                existing.save()
                updated_count += 1
            else:
                RadioLabInventory.objects.create(
                    item=item_name,
                    item_id=proc_code,
                    rate=rate_val,
                    type='R',
                    sponsor=target_sponsor,
                    plan=target_plan,
                    staff=staff
                )
                created_count += 1

        msg = f"Bulk upload completed: {created_count} radiology tariffs created, {updated_count} updated."
        if skipped:
            msg += f" {len(skipped)} rows were skipped."

        return Response({
            'detail': msg,
            'created': created_count,
            'updated': updated_count,
            'skipped': skipped,
            'total_processed': created_count + updated_count + len(skipped),
        }, status=status.HTTP_200_OK)


# 1.4 Admission Fees (Wards)
class AdmissionFeesView(APIView):
    permission_classes = [permissions.AllowAny]

    def get(self, request):
        wards = Ward.objects.all().order_by('ward_name')
        data = [{'id': w.id, 'ward_name': w.ward_name, 'price': w.price} for w in wards]
        return Response({'total': len(data), 'wards': data})

    def post(self, request):
        ward_name = request.data.get('ward_name', '').strip()
        price = request.data.get('price', 0)

        if not ward_name:
            return Response({'detail': 'Ward name required.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            price = int(price)
            if price < 0:
                return Response({'detail': 'Price cannot be negative.'}, status=status.HTTP_400_BAD_REQUEST)
        except (ValueError, TypeError):
            return Response({'detail': 'Invalid price format.'}, status=status.HTTP_400_BAD_REQUEST)

        if Ward.objects.filter(ward_name__iexact=ward_name).exists():
            return Response({'detail': f'Ward "{ward_name}" already exists.'}, status=status.HTTP_400_BAD_REQUEST)

        w = Ward.objects.create(ward_name=ward_name, price=price)
        return Response({'detail': f'Ward "{w.ward_name}" created.', 'id': w.id, 'ward_name': w.ward_name, 'price': w.price})

    def put(self, request, pk=None):
        try:
            w = Ward.objects.get(id=pk)
        except Ward.DoesNotExist:
            return Response({'detail': 'Ward not found.'}, status=status.HTTP_404_NOT_FOUND)

        ward_name = request.data.get('ward_name', w.ward_name).strip()
        price = request.data.get('price', w.price)

        if Ward.objects.filter(ward_name__iexact=ward_name).exclude(id=pk).exists():
            return Response({'detail': f'Ward with name "{ward_name}" already exists.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            w.price = int(price)
        except (ValueError, TypeError):
            return Response({'detail': 'Invalid price format.'}, status=status.HTTP_400_BAD_REQUEST)

        w.ward_name = ward_name
        w.save()
        return Response({'detail': 'Ward updated.', 'id': w.id, 'ward_name': w.ward_name, 'price': w.price})

    def delete(self, request, pk=None):
        try:
            Ward.objects.filter(id=pk).delete()
            return Response({'detail': 'Ward deleted.'})
        except Exception as e:
            return Response({'detail': str(e)}, status=status.HTTP_400_BAD_REQUEST)


# 1.5 Other Services Fees
class OtherServicesView(APIView):
    permission_classes = [permissions.AllowAny]

    def get(self, request):
        qs = OtherService2.objects.all().select_related('plan').order_by('service')
        plan_id = request.query_params.get('plan')
        patient_id = request.query_params.get('patient')
        search = request.query_params.get('search', '').strip()

        if patient_id:
            try:
                patient = PatientProfile.objects.select_related('plan').get(id=patient_id)
                if patient.plan:
                    qs = qs.filter(Q(plan=patient.plan) | Q(plan__isnull=True))
                else:
                    qs = qs.filter(plan__isnull=True)
            except PatientProfile.DoesNotExist:
                pass
        elif plan_id and str(plan_id).lower() != 'all':
            qs = qs.filter(Q(plan_id=plan_id) | Q(plan__isnull=True))

        if search:
            qs = qs.filter(Q(service__icontains=search) | Q(service_id__icontains=search))

        data = []
        for s in qs:
            data.append({
                'id': s.id,
                'service': s.service,
                'service_id': s.service_id or '',
                'rate': float(s.rate),
                'plan_id': s.plan.id if s.plan else None,
                'plan_name': s.plan.plan if s.plan else 'All Plans',
            })
        return Response({'total': len(data), 'services': data, 'results': data})

    def post(self, request):
        service = request.data.get('service', '').strip()
        service_id = request.data.get('service_id', '').strip()
        rate = request.data.get('rate', 0)
        plan_id = request.data.get('plan_id')

        if not service:
            return Response({'detail': 'Service name required.'}, status=status.HTTP_400_BAD_REQUEST)

        plan = None
        if plan_id:
            try:
                plan = PatientPlan.objects.get(id=plan_id)
            except PatientPlan.DoesNotExist:
                pass

        staff = request.user if request.user.is_authenticated else User.objects.first()

        obj = OtherService2.objects.create(
            service=service,
            service_id=service_id or f"SRV-{OtherService2.objects.count()+1:03d}",
            rate=Decimal(str(rate)),
            plan=plan,
            staff=staff
        )
        return Response({
            'detail': f'Service "{obj.service}" created.',
            'id': obj.id,
            'service': obj.service,
            'service_id': obj.service_id,
            'rate': float(obj.rate),
            'plan_name': obj.plan.plan if obj.plan else 'All Plans'
        })

    def put(self, request, pk=None):
        try:
            obj = OtherService2.objects.get(id=pk)
        except OtherService2.DoesNotExist:
            return Response({'detail': 'Service not found.'}, status=status.HTTP_404_NOT_FOUND)

        if 'service' in request.data: obj.service = request.data['service'].strip()
        if 'rate' in request.data: obj.rate = Decimal(str(request.data['rate']))
        if 'service_id' in request.data: obj.service_id = request.data['service_id'].strip()

        obj.save()
        return Response({'detail': 'Service updated.', 'id': obj.id, 'service': obj.service, 'rate': float(obj.rate)})

    def delete(self, request, pk=None):
        OtherService2.objects.filter(id=pk).delete()
        return Response({'detail': 'Service deleted.'})


# ==========================================================
# 2. PRODUCT TARIFFS
# ==========================================================

# 2.1 Medication Fees (Pharmacy Tariff)
class MedicationFeesView(APIView):
    permission_classes = [permissions.AllowAny]

    def get(self, request):
        qs = PharmacyTariff.objects.all().select_related('sponsor', 'plan', 'plan__sponsor').order_by('product_name')

        tariff_type = request.query_params.get('tariff_type', '').strip().lower()
        patient_id = request.query_params.get('patient')
        sponsor_id = request.query_params.get('sponsor')
        plan_id = request.query_params.get('plan')
        search = request.query_params.get('search', '').strip()

        # If explicitly requesting general out-of-pocket tariffs
        if tariff_type == 'general' or (sponsor_id and str(sponsor_id).lower() in ('general', 'self', 'cash')):
            qs = qs.filter(sponsor__isnull=True, plan__isnull=True)
        # If filtering for a specific patient
        elif patient_id or tariff_type == 'covered':
            if patient_id:
                try:
                    patient = PatientProfile.objects.select_related('sponsor', 'plan', 'category', 'sponsor__category').get(id=patient_id)
                    patient_sponsor = patient.sponsor
                    patient_plan = patient.plan

                    # Check if patient or sponsor has a Private category
                    is_private_patient = (
                        (patient.category and 'private' in str(patient.category.category).lower()) or
                        (patient_sponsor and patient_sponsor.category and 'private' in str(patient_sponsor.category.category).lower()) or
                        (patient_sponsor and ('private' in patient_sponsor.name.lower() or 'self' in patient_sponsor.name.lower())) or
                        (not patient_sponsor and not patient_plan)
                    )

                    if is_private_patient:
                        # Private sponsor maintains the same price as general!
                        # General price is always available for them because they pay out of pocket
                        qs = qs.filter(
                            (Q(sponsor__isnull=True) & Q(plan__isnull=True)) |
                            (Q(sponsor=patient_sponsor) if patient_sponsor else Q(pk__in=[]))
                        )
                    else:
                        # HMO Sponsor: uses covered plan tariffs (claims requiring HMO approval)
                        if not patient_sponsor and sponsor_id and str(sponsor_id).lower() not in ('all', 'general'):
                            try:
                                patient_sponsor = Sponsor.objects.get(id=sponsor_id)
                            except Sponsor.DoesNotExist:
                                pass
                        if not patient_plan and plan_id and str(plan_id).lower() not in ('all', 'general'):
                            try:
                                patient_plan = PatientPlan.objects.get(id=plan_id)
                            except PatientPlan.DoesNotExist:
                                pass

                        if tariff_type == 'general':
                            qs = qs.filter(sponsor__isnull=True, plan__isnull=True)
                        elif patient_sponsor and patient_plan:
                            qs = qs.filter(
                                (Q(sponsor=patient_sponsor) & Q(plan=patient_plan)) |
                                Q(plan=patient_plan) |
                                (Q(sponsor=patient_sponsor) & Q(plan__isnull=True))
                            )
                        elif patient_plan:
                            qs = qs.filter(plan=patient_plan)
                        elif patient_sponsor:
                            qs = qs.filter(Q(sponsor=patient_sponsor) | Q(plan__sponsor=patient_sponsor))
                        else:
                            qs = qs.filter(sponsor__isnull=True, plan__isnull=True)
                except PatientProfile.DoesNotExist:
                    if sponsor_id and str(sponsor_id).lower() != 'all':
                        qs = qs.filter(Q(sponsor_id=sponsor_id) | Q(plan__sponsor_id=sponsor_id))
                    if plan_id and str(plan_id).lower() != 'all':
                        qs = qs.filter(plan_id=plan_id)
            else:
                if sponsor_id and str(sponsor_id).lower() != 'all':
                    qs = qs.filter(Q(sponsor_id=sponsor_id) | Q(plan__sponsor_id=sponsor_id))
                if plan_id and str(plan_id).lower() != 'all':
                    qs = qs.filter(plan_id=plan_id)
        else:
            if sponsor_id and str(sponsor_id).lower() != 'all':
                qs = qs.filter(Q(sponsor_id=sponsor_id) | Q(plan__sponsor_id=sponsor_id))
            if plan_id and str(plan_id).lower() != 'all':
                qs = qs.filter(plan_id=plan_id)

        if search:
            qs = qs.filter(
                Q(product_name__icontains=search) |
                Q(product_id__icontains=search)
            )

        data = []
        for p in qs:
            eff_sponsor = p.sponsor or (p.plan.sponsor if p.plan else None)
            is_gen = bool(not eff_sponsor and not p.plan)
            data.append({
                'id': p.id,
                'product_name': p.product_name,
                'product_id': p.product_id,
                'sku': p.product_id,
                'rate': float(p.rate),
                'sponsor_id': eff_sponsor.id if eff_sponsor else None,
                'sponsor_name': eff_sponsor.name if eff_sponsor else 'General / Self-Pay',
                'sponsor_code': eff_sponsor.code if eff_sponsor else 'GENERAL',
                'plan_id': p.plan.id if p.plan else None,
                'plan_name': p.plan.plan if p.plan else ('General Price' if is_gen else 'All Plans'),
                'plan_code': p.plan.code if p.plan else ('GENERAL' if is_gen else ''),
                'is_general': is_gen,
                'payment_label': 'Pay from your purse' if is_gen else 'Covered by Plan',
            })
        return Response({'total': len(data), 'medications': data, 'results': data})

    def post(self, request):
        product_name = request.data.get('product_name', '').strip()
        product_id = request.data.get('product_id', '').strip()
        rate = request.data.get('rate', 0)
        sponsor_id = request.data.get('sponsor_id') or request.data.get('sponsor')
        plan_id = request.data.get('plan_id') or request.data.get('plan')

        if not product_name:
            return Response({'detail': 'Product name required.'}, status=status.HTTP_400_BAD_REQUEST)

        is_general = (
            not sponsor_id or
            str(sponsor_id).lower() in ('', 'general', 'self', 'cash', 'none', 'null')
        )

        plan = None
        sponsor = None

        if not is_general:
            if plan_id and str(plan_id).lower() not in ('all', 'general', 'none'):
                try:
                    plan = PatientPlan.objects.select_related('sponsor').get(id=plan_id)
                except PatientPlan.DoesNotExist:
                    return Response({'detail': 'Selected Plan not found.'}, status=status.HTTP_404_NOT_FOUND)

            if sponsor_id and str(sponsor_id).lower() not in ('all', 'general', 'none'):
                try:
                    sponsor = Sponsor.objects.get(id=sponsor_id)
                except Sponsor.DoesNotExist:
                    return Response({'detail': 'Selected Sponsor not found.'}, status=status.HTTP_404_NOT_FOUND)
            elif plan and plan.sponsor:
                sponsor = plan.sponsor

        staff = request.user if request.user.is_authenticated else User.objects.first()

        obj = PharmacyTariff.objects.create(
            product_name=product_name,
            product_id=product_id or f"DRG-{PharmacyTariff.objects.count()+1:03d}",
            rate=Decimal(str(rate)),
            sponsor=sponsor,
            plan=plan,
            staff=staff
        )

        eff_sponsor = obj.sponsor or (obj.plan.sponsor if obj.plan else None)
        is_gen = bool(not eff_sponsor and not obj.plan)
        return Response({
            'detail': f'Medication tariff "{obj.product_name}" created successfully.',
            'id': obj.id,
            'product_name': obj.product_name,
            'product_id': obj.product_id,
            'rate': float(obj.rate),
            'sponsor_id': eff_sponsor.id if eff_sponsor else None,
            'sponsor_name': eff_sponsor.name if eff_sponsor else 'General / Self-Pay',
            'sponsor_code': eff_sponsor.code if eff_sponsor else 'GENERAL',
            'plan_id': obj.plan.id if obj.plan else None,
            'plan_name': obj.plan.plan if obj.plan else ('General Price' if is_gen else 'All Plans'),
            'is_general': is_gen,
            'payment_label': 'Pay from your purse' if is_gen else 'Covered by Plan',
        }, status=status.HTTP_201_CREATED)

    def put(self, request, pk=None):
        try:
            obj = PharmacyTariff.objects.select_related('sponsor', 'plan').get(id=pk)
        except PharmacyTariff.DoesNotExist:
            return Response({'detail': 'Medication tariff not found.'}, status=status.HTTP_404_NOT_FOUND)

        if 'product_name' in request.data:
            obj.product_name = request.data['product_name'].strip()
        if 'rate' in request.data:
            obj.rate = Decimal(str(request.data['rate']))
        if 'product_id' in request.data:
            obj.product_id = request.data['product_id'].strip()

        sponsor_id = request.data.get('sponsor_id') or request.data.get('sponsor')
        if sponsor_id is not None:
            if sponsor_id and str(sponsor_id).lower() not in ('all', 'general', 'self', 'none', 'null'):
                try:
                    obj.sponsor = Sponsor.objects.get(id=sponsor_id)
                except Sponsor.DoesNotExist:
                    pass
            else:
                obj.sponsor = None

        plan_id = request.data.get('plan_id') or request.data.get('plan')
        if plan_id is not None:
            if plan_id and str(plan_id).lower() not in ('all', 'general', 'none', 'null'):
                try:
                    obj.plan = PatientPlan.objects.get(id=plan_id)
                    if obj.plan.sponsor and not obj.sponsor:
                        obj.sponsor = obj.plan.sponsor
                except PatientPlan.DoesNotExist:
                    pass
            else:
                obj.plan = None

        obj.save()
        eff_sponsor = obj.sponsor or (obj.plan.sponsor if obj.plan else None)
        is_gen = bool(not eff_sponsor and not obj.plan)
        return Response({
            'detail': 'Medication tariff updated successfully.',
            'id': obj.id,
            'product_name': obj.product_name,
            'product_id': obj.product_id,
            'rate': float(obj.rate),
            'sponsor_id': eff_sponsor.id if eff_sponsor else None,
            'sponsor_name': eff_sponsor.name if eff_sponsor else 'General / Self-Pay',
            'plan_id': obj.plan.id if obj.plan else None,
            'plan_name': obj.plan.plan if obj.plan else ('General Price' if is_gen else 'All Plans'),
            'is_general': is_gen,
            'payment_label': 'Pay from your purse' if is_gen else 'Covered by Plan',
        })

    def delete(self, request, pk=None):
        PharmacyTariff.objects.filter(id=pk).delete()
        return Response({'detail': 'Medication tariff deleted.'})


# 2.1b Load / Sync General Medication Prices (Hospital Formulary / Inventory Standard Prices)
class LoadGeneralMedicationPricesView(APIView):
    permission_classes = [permissions.AllowAny]

    def post(self, request):
        from inventory.models import Product, PharmacyTariff
        from django.db.models import Max

        staff = request.user if request.user.is_authenticated else User.objects.first()
        created_count = 0
        updated_count = 0

        # Baseline clinical formulary prices for general self-pay patients
        BASELINE_MED_PRICES = {
            'Paracetamol 500mg Tablet': ('MED-001', Decimal('500.00')),
            'Amoxicillin 500mg Capsule': ('MED-002', Decimal('1200.00')),
            'Artemether/Lumefantrine 20/120mg (Coartem)': ('MED-003', Decimal('2500.00')),
            'Ciprofloxacin 500mg Tablet': ('MED-004', Decimal('1800.00')),
            'Metformin 500mg Tablet': ('MED-005', Decimal('950.00')),
            'Amlodipine 5mg Tablet': ('MED-006', Decimal('1400.00')),
            'Ibuprofen 400mg Tablet': ('MED-007', Decimal('600.00')),
            'Omeprazole 20mg Capsule': ('MED-008', Decimal('1100.00')),
            'Ceftriaxone 1g Injection': ('MED-009', Decimal('3500.00')),
            'Oral Rehydration Salts (ORS) Sachet': ('MED-010', Decimal('300.00')),
            'Dextromethorphan Syrup': ('MED-011', Decimal('1500.00')),
        }

        # 1. Sync from inventory Product model if any exist
        inv_products = Product.objects.filter(price__gt=0)
        for prod in inv_products:
            sku = prod.product_id or f"SKU-{prod.id:04d}"
            tariff, created = PharmacyTariff.objects.get_or_create(
                product_name=prod.product_name,
                sponsor__isnull=True,
                plan__isnull=True,
                defaults={
                    'product_id': sku,
                    'rate': prod.price,
                    'staff': staff,
                    'sponsor': None,
                    'plan': None,
                }
            )
            if created:
                created_count += 1
            else:
                tariff.rate = prod.price
                tariff.product_id = sku
                tariff.save()
                updated_count += 1

        # 2. Sync all distinct products from PharmacyTariff that don't have general rate yet
        existing_names = list(PharmacyTariff.objects.values_list('product_name', flat=True).distinct())
        all_drug_names = set(existing_names).union(BASELINE_MED_PRICES.keys())

        for name in all_drug_names:
            if not name:
                continue
            gen_tariff = PharmacyTariff.objects.filter(
                product_name__iexact=name.strip(),
                sponsor__isnull=True,
                plan__isnull=True
            ).first()

            if not gen_tariff:
                baseline = BASELINE_MED_PRICES.get(name.strip())
                if baseline:
                    sku, rate = baseline
                else:
                    existing_t = PharmacyTariff.objects.filter(product_name__iexact=name.strip()).first()
                    sku = (existing_t.product_id if existing_t else None) or f"GEN-{PharmacyTariff.objects.count()+1:03d}"
                    max_rate = PharmacyTariff.objects.filter(product_name__iexact=name.strip()).aggregate(Max('rate'))['rate__max']
                    rate = max_rate or Decimal('1000.00')

                PharmacyTariff.objects.create(
                    product_name=name.strip(),
                    product_id=sku,
                    rate=rate,
                    sponsor=None,
                    plan=None,
                    staff=staff
                )
                created_count += 1

                # Keep Product inventory table in sync
                Product.objects.get_or_create(
                    product_name=name.strip(),
                    defaults={
                        'product_id': sku,
                        'price': rate,
                        'stock': 100,
                        'staff': staff,
                    }
                )

        total_general = PharmacyTariff.objects.filter(sponsor__isnull=True, plan__isnull=True).count()
        return Response({
            'detail': f'Loaded {created_count} new general medication tariffs ({total_general} total active general rates).',
            'created': created_count,
            'updated': updated_count,
            'total_general': total_general,
        })


class LoadGeneralLabPricesView(APIView):
    permission_classes = [permissions.AllowAny]

    def post(self, request):
        from radio_lab.models import RadioLabInventory
        from django.db.models import Max

        staff = request.user if request.user.is_authenticated else User.objects.first()
        created_count = 0
        updated_count = 0

        # Baseline clinical laboratory investigation prices for general self-pay patients
        BASELINE_LAB_PRICES = {
            'Full Blood Count (FBC / CBC)': ('LAB-001', Decimal('3000.00')),
            'Malaria Parasite (MP by Microscopy)': ('LAB-002', Decimal('1500.00')),
            'Widal Agglutination Test (Typhoid)': ('LAB-003', Decimal('2000.00')),
            'Urinalysis (10-Parameter Strip)': ('LAB-004', Decimal('1200.00')),
            'Fasting Blood Sugar (FBS)': ('LAB-005', Decimal('1000.00')),
            'Lipid Profile (Full)': ('LAB-006', Decimal('6000.00')),
            'Liver Function Test (LFT)': ('LAB-007', Decimal('5000.00')),
            'Kidney Function Test (Electrolytes, Urea, Creatinine)': ('LAB-008', Decimal('5500.00')),
            'Stool Microscopy, Culture & Sensitivity': ('LAB-009', Decimal('3500.00')),
            'Hepatitis B Surface Antigen (HBsAg Screening)': ('LAB-010', Decimal('2500.00')),
        }

        # Sync all distinct laboratory tests from existing RadioLabInventory (type='L')
        existing_items = list(RadioLabInventory.objects.filter(type='L').values_list('item', flat=True).distinct())
        all_lab_names = set(existing_items).union(BASELINE_LAB_PRICES.keys())

        for name in all_lab_names:
            if not name:
                continue
            gen_tariff = RadioLabInventory.objects.filter(
                item__iexact=name.strip(),
                type='L',
                sponsor__isnull=True,
                plan__isnull=True
            ).first()

            if not gen_tariff:
                baseline = BASELINE_LAB_PRICES.get(name.strip())
                if baseline:
                    code, rate = baseline
                else:
                    existing_t = RadioLabInventory.objects.filter(item__iexact=name.strip(), type='L').first()
                    code = (existing_t.item_id if existing_t else None) or f"LAB-{RadioLabInventory.objects.filter(type='L').count()+1:03d}"
                    max_rate = RadioLabInventory.objects.filter(item__iexact=name.strip(), type='L').aggregate(Max('rate'))['rate__max']
                    rate = max_rate or Decimal('2000.00')

                RadioLabInventory.objects.create(
                    item=name.strip(),
                    item_id=code,
                    rate=rate,
                    type='L',
                    sponsor=None,
                    plan=None,
                    staff=staff
                )
                created_count += 1
            else:
                updated_count += 1

        total_general = RadioLabInventory.objects.filter(type='L', sponsor__isnull=True, plan__isnull=True).count()
        return Response({
            'detail': f'Loaded {created_count} general laboratory tariffs ({total_general} total active general rates).',
            'created': created_count,
            'updated': updated_count,
            'total_general': total_general,
        })


# 2.2 Download Medication Tariff Template
class DownloadMedicationTariffTemplateView(APIView):
    permission_classes = [permissions.AllowAny]

    def get(self, request):
        from django.http import HttpResponse
        import io, csv

        sponsor_id = request.query_params.get('sponsor')
        plan_id = request.query_params.get('plan')

        sponsor_code = "AMH"
        plan_code = "AIICO SLIVER"
        sponsor_name = "AIICO MULTISHIELD (AMH)"
        plan_name = "AIICO SLIVER"

        if sponsor_id:
            try:
                s = Sponsor.objects.get(id=sponsor_id)
                sponsor_code = s.code or s.name
                sponsor_name = s.name
            except Sponsor.DoesNotExist:
                pass

        if plan_id:
            try:
                p = PatientPlan.objects.select_related('sponsor').get(id=plan_id)
                plan_code = p.code or p.plan
                plan_name = p.plan
                if p.sponsor:
                    sponsor_code = p.sponsor.code or p.sponsor.name
                    sponsor_name = p.sponsor.name
            except PatientPlan.DoesNotExist:
                pass

        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(['product_name', 'sku', 'rate', 'sponsor_code', 'plan_code', 'notes'])

        samples = [
            ['Paracetamol 500mg Tablet', 'MED-001', '500.00', sponsor_code, plan_code, f'Covered under {plan_name}'],
            ['Amoxicillin 500mg Capsule', 'MED-002', '1200.00', sponsor_code, plan_code, f'Covered under {plan_name}'],
            ['Artemether/Lumefantrine 20/120mg (Coartem)', 'MED-003', '2500.00', sponsor_code, plan_code, f'Covered under {plan_name}'],
            ['Ciprofloxacin 500mg Tablet', 'MED-004', '1800.00', sponsor_code, plan_code, f'Covered under {plan_name}'],
            ['Metformin 500mg Tablet', 'MED-005', '950.00', sponsor_code, plan_code, f'Covered under {plan_name}'],
            ['Amlodipine 5mg Tablet', 'MED-006', '1400.00', sponsor_code, plan_code, f'Covered under {plan_name}'],
            ['Ibuprofen 400mg Tablet', 'MED-007', '600.00', sponsor_code, plan_code, f'Covered under {plan_name}'],
            ['Omeprazole 20mg Capsule', 'MED-008', '1100.00', sponsor_code, plan_code, f'Covered under {plan_name}'],
            ['Ceftriaxone 1g Injection', 'MED-009', '3500.00', sponsor_code, plan_code, f'Covered under {plan_name}'],
            ['Oral Rehydration Salts (ORS) Sachet', 'MED-010', '300.00', sponsor_code, plan_code, f'Covered under {plan_name}'],
        ]
        for row in samples:
            writer.writerow(row)

        csv_content = output.getvalue()
        response = HttpResponse(csv_content, content_type='text/csv; charset=utf-8')
        safe_sponsor = str(sponsor_code).replace(' ', '_').replace('/', '_')
        safe_plan = str(plan_code).replace(' ', '_').replace('/', '_')
        filename = f"medication_tariffs_template_{safe_sponsor}_{safe_plan}.csv"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


# 2.3 Bulk Upload Medication Tariffs
class BulkUploadMedicationTariffView(APIView):
    permission_classes = [permissions.AllowAny]

    def post(self, request):
        import io, csv
        default_sponsor_id = request.data.get('sponsor_id') or request.data.get('default_sponsor_id')
        default_plan_id = request.data.get('plan_id') or request.data.get('default_plan_id')

        default_sponsor = None
        if default_sponsor_id and str(default_sponsor_id).lower() not in ('', 'all', 'null', 'none'):
            try:
                default_sponsor = Sponsor.objects.get(id=default_sponsor_id)
            except Sponsor.DoesNotExist:
                pass

        default_plan = None
        if default_plan_id and str(default_plan_id).lower() not in ('', 'all', 'null', 'none'):
            try:
                default_plan = PatientPlan.objects.select_related('sponsor').get(id=default_plan_id)
                if not default_sponsor and default_plan.sponsor:
                    default_sponsor = default_plan.sponsor
            except PatientPlan.DoesNotExist:
                pass

        raw_rows = []
        uploaded_file = request.FILES.get('file')

        if uploaded_file:
            filename = uploaded_file.name.lower()
            try:
                if filename.endswith('.csv'):
                    decoded_file = uploaded_file.read().decode('utf-8-sig', errors='ignore')
                    reader = csv.DictReader(io.StringIO(decoded_file))
                    for row in reader:
                        raw_rows.append(row)
                elif filename.endswith(('.xlsx', '.xls')):
                    import pandas as pd
                    df = pd.read_excel(io.BytesIO(uploaded_file.read()))
                    df = df.where(pd.notnull(df), None)
                    raw_rows = df.to_dict(orient='records')
                else:
                    return Response(
                        {'detail': 'Unsupported file format. Please upload a CSV (.csv) or Excel (.xlsx, .xls) file.'},
                        status=status.HTTP_400_BAD_REQUEST
                    )
            except Exception as e:
                return Response(
                    {'detail': f'Error reading uploaded file: {str(e)}'},
                    status=status.HTTP_400_BAD_REQUEST
                )
        elif 'items' in request.data:
            items = request.data.get('items', [])
            if isinstance(items, list):
                raw_rows = items
        else:
            return Response(
                {'detail': 'No file or items array provided for bulk upload.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        if not raw_rows:
            return Response(
                {'detail': 'The uploaded file contains no data rows.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        # Cache sponsors and plans in memory for fast lookup
        all_sponsors = list(Sponsor.objects.all())
        all_plans = list(PatientPlan.objects.select_related('sponsor').all())

        sponsor_by_code = {s.code.strip().upper(): s for s in all_sponsors if s.code}
        sponsor_by_name = {s.name.strip().upper(): s for s in all_sponsors if s.name}
        sponsor_by_id = {str(s.id): s for s in all_sponsors}

        plan_by_code = {p.code.strip().upper(): p for p in all_plans if p.code}
        plan_by_name = {p.plan.strip().upper(): p for p in all_plans if p.plan}
        plan_by_id = {str(p.id): p for p in all_plans}

        staff = request.user if request.user.is_authenticated else User.objects.first()

        created_count = 0
        updated_count = 0
        skipped = []

        counter = PharmacyTariff.objects.count()

        for idx, row in enumerate(raw_rows, start=1):
            if not isinstance(row, dict):
                continue

            clean_row = {str(k).strip().lower(): v for k, v in row.items() if k is not None}

            product_name = (
                clean_row.get('product_name')
                or clean_row.get('medication_name')
                or clean_row.get('medication')
                or clean_row.get('drug_name')
                or clean_row.get('drug')
                or clean_row.get('name')
                or clean_row.get('item')
            )
            if not product_name or not str(product_name).strip():
                skipped.append(f"Row {idx}: Missing product name")
                continue

            product_name = str(product_name).strip()

            rate_raw = (
                clean_row.get('rate')
                or clean_row.get('tariff_rate')
                or clean_row.get('price')
                or clean_row.get('amount')
                or clean_row.get('tariff')
            )
            try:
                rate_val = Decimal(str(rate_raw).replace(',', '').replace('₦', '').strip())
                if rate_val < 0:
                    rate_val = Decimal('0.00')
            except Exception:
                skipped.append(f"Row {idx} ({product_name}): Invalid rate '{rate_raw}'")
                continue

            sku = (
                clean_row.get('sku')
                or clean_row.get('product_id')
                or clean_row.get('code')
                or clean_row.get('item_code')
            )
            if sku:
                sku = str(sku).strip()[:13]
            else:
                counter += 1
                sku = f"MED-{counter:04d}"

            # Resolve Plan
            target_plan = default_plan
            plan_hint = clean_row.get('plan_code') or clean_row.get('plan') or clean_row.get('plan_name') or clean_row.get('coverage_plan')
            if plan_hint:
                plan_key = str(plan_hint).strip().upper()
                if plan_key in plan_by_code:
                    target_plan = plan_by_code[plan_key]
                elif plan_key in plan_by_name:
                    target_plan = plan_by_name[plan_key]
                elif str(plan_hint).strip() in plan_by_id:
                    target_plan = plan_by_id[str(plan_hint).strip()]
                else:
                    # Fuzzy match fallback (e.g. "Silver", "Gold", "Bronze")
                    clean_hint = plan_key.replace(' ', '').replace('_', '')
                    for p in all_plans:
                        p_clean = (p.plan or '').upper().replace(' ', '').replace('_', '')
                        if clean_hint and (clean_hint in p_clean or p_clean in clean_hint):
                            target_plan = p
                            break
                        if ('SILVER' in clean_hint and 'SLIVER' in p_clean) or ('SLIVER' in clean_hint and 'SILVER' in p_clean):
                            target_plan = p
                            break

            # Resolve Sponsor
            target_sponsor = default_sponsor
            sponsor_hint = clean_row.get('sponsor_code') or clean_row.get('sponsor') or clean_row.get('sponsor_name') or clean_row.get('hmo')
            if sponsor_hint:
                s_key = str(sponsor_hint).strip().upper()
                if s_key in ('GENERAL', 'SELF', 'SELF-PAY', 'CASH', 'OUT-OF-POCKET', 'NO HMO', 'NONE'):
                    target_sponsor = None
                    target_plan = None
                elif s_key in sponsor_by_code:
                    target_sponsor = sponsor_by_code[s_key]
                elif s_key in sponsor_by_name:
                    target_sponsor = sponsor_by_name[s_key]
                elif str(sponsor_hint).strip() in sponsor_by_id:
                    target_sponsor = sponsor_by_id[str(sponsor_hint).strip()]

            if not target_sponsor and target_plan and target_plan.sponsor:
                target_sponsor = target_plan.sponsor

            # Look for existing tariff for this product and (sponsor, plan)
            existing = PharmacyTariff.objects.filter(
                Q(product_name__iexact=product_name) | Q(product_id__iexact=sku),
                sponsor=target_sponsor,
                plan=target_plan
            ).first()

            if existing:
                existing.rate = rate_val
                existing.product_name = product_name
                existing.product_id = sku
                existing.staff = staff
                existing.save()
                updated_count += 1
            else:
                PharmacyTariff.objects.create(
                    product_name=product_name,
                    product_id=sku,
                    rate=rate_val,
                    sponsor=target_sponsor,
                    plan=target_plan,
                    staff=staff
                )
                created_count += 1

        total_processed = created_count + updated_count
        return Response({
            'detail': f'Bulk import finished: {created_count} added, {updated_count} updated, {len(skipped)} skipped.',
            'created_count': created_count,
            'updated_count': updated_count,
            'skipped_count': len(skipped),
            'total_processed': total_processed,
            'skipped_details': skipped[:10],
        }, status=status.HTTP_200_OK if total_processed > 0 else status.HTTP_400_BAD_REQUEST)


# ==========================================================
# 3. MANAGE PACKAGES
# ==========================================================

# 3.1 Packages
class PackagesView(APIView):
    permission_classes = [permissions.AllowAny]

    def get(self, request):
        packages = Packages.objects.all().select_related('staff').order_by('-created_date')
        data = []
        for p in packages:
            data.append({
                'id': p.id,
                'name': p.name,
                'staff_name': p.staff.fullname if p.staff else 'Admin',
                'created_date': p.created_date.isoformat() if p.created_date else None,
                'items_count': PackagesData.objects.filter(package=p).count()
            })
        return Response({'total': len(data), 'packages': data})

    def post(self, request):
        name = request.data.get('name', '').strip()
        if not name:
            return Response({'detail': 'Package name required.'}, status=status.HTTP_400_BAD_REQUEST)

        if Packages.objects.filter(name__iexact=name).exists():
            return Response({'detail': f'Package "{name}" already exists.'}, status=status.HTTP_400_BAD_REQUEST)

        staff = request.user if request.user.is_authenticated else User.objects.first()
        pkg = Packages.objects.create(name=name, staff=staff)
        return Response({'detail': f'Package "{pkg.name}" created.', 'id': pkg.id, 'name': pkg.name})

    def put(self, request, pk=None):
        try:
            pkg = Packages.objects.get(id=pk)
        except Packages.DoesNotExist:
            return Response({'detail': 'Package not found.'}, status=status.HTTP_404_NOT_FOUND)

        name = request.data.get('name', '').strip()
        if not name:
            return Response({'detail': 'Package name required.'}, status=status.HTTP_400_BAD_REQUEST)

        pkg.name = name
        pkg.save()
        return Response({'detail': 'Package updated.', 'id': pkg.id, 'name': pkg.name})

    def delete(self, request, pk=None):
        Packages.objects.filter(id=pk).delete()
        return Response({'detail': 'Package deleted.'})


# 3.2 Package Data Items
class PackageDataView(APIView):
    permission_classes = [permissions.AllowAny]

    def get(self, request):
        package_id = request.query_params.get('package_id')
        qs = PackagesData.objects.all().select_related('package').order_by('type')
        if package_id:
            qs = qs.filter(package_id=package_id)

        data = []
        for item in qs:
            data.append({
                'id': item.id,
                'package_id': item.package.id if item.package else None,
                'package_name': item.package.name if item.package else '',
                'item': item.item,
                'type': item.type,
                'rate': float(item.rate),
            })
        return Response({'total': len(data), 'items': data})

    def post(self, request):
        package_id = request.data.get('package_id')
        item_name = request.data.get('item', '').strip()
        item_type = request.data.get('type', 'SR').strip()
        rate = request.data.get('rate', 0)

        if not package_id or not item_name:
            return Response({'detail': 'Package ID and item name are required.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            pkg = Packages.objects.get(id=package_id)
        except Packages.DoesNotExist:
            return Response({'detail': 'Package not found.'}, status=status.HTTP_404_NOT_FOUND)

        staff = request.user if request.user.is_authenticated else User.objects.first()

        obj = PackagesData.objects.create(
            package=pkg,
            item=item_name,
            type=item_type,
            rate=Decimal(str(rate)),
            staff=staff
        )
        return Response({
            'detail': f'Item "{obj.item}" added to package "{pkg.name}".',
            'id': obj.id,
            'package_id': pkg.id,
            'item': obj.item,
            'type': obj.type,
            'rate': float(obj.rate)
        })

    def put(self, request, pk=None):
        try:
            item = PackagesData.objects.get(id=pk)
        except PackagesData.DoesNotExist:
            return Response({'detail': 'Package item not found.'}, status=status.HTTP_404_NOT_FOUND)

        if 'item' in request.data: item.item = request.data['item'].strip()
        if 'type' in request.data: item.type = request.data['type'].strip()
        if 'rate' in request.data: item.rate = Decimal(str(request.data['rate']))

        item.save()
        return Response({'detail': 'Package item updated.', 'id': item.id, 'item': item.item, 'rate': float(item.rate)})

    def delete(self, request, pk=None):
        PackagesData.objects.filter(id=pk).delete()
        return Response({'detail': 'Package item deleted.'})


# 3.3 Immunization Fees
class ImmunizationFeesView(APIView):
    permission_classes = [permissions.AllowAny]

    def get(self, request):
        # Sample immunization catalog fees
        immunizations = [
            {'id': 1, 'vaccine_name': 'BCG (Bacille Calmette-Guérin)', 'target': 'Newborns', 'schedule': 'At Birth', 'fee': 3500},
            {'id': 2, 'vaccine_name': 'OPV (Oral Polio Vaccine)', 'target': 'Infants', 'schedule': 'Birth, 6, 10, 14 Weeks', 'fee': 2500},
            {'id': 3, 'vaccine_name': 'Pentavalent (DTP-HepB-Hib)', 'target': 'Infants', 'schedule': '6, 10, 14 Weeks', 'fee': 6000},
            {'id': 4, 'vaccine_name': 'Rotavirus Vaccine', 'target': 'Infants', 'schedule': '6, 10 Weeks', 'fee': 8500},
            {'id': 5, 'vaccine_name': 'Yellow Fever Vaccine', 'target': 'Infants / Adults', 'schedule': '9 Months / Booster', 'fee': 5000},
            {'id': 6, 'vaccine_name': 'Measles Vaccine', 'target': 'Infants', 'schedule': '9 Months', 'fee': 4000},
            {'id': 7, 'vaccine_name': 'Tetanus Toxoid (TT)', 'target': 'Maternal / Adults', 'schedule': 'Antenatal Schedule', 'fee': 2000},
            {'id': 8, 'vaccine_name': 'Hepatitis B Adult Booster', 'target': 'Adults / Staff', 'schedule': '0, 1, 6 Months', 'fee': 7500},
        ]
        return Response({'total': len(immunizations), 'immunizations': immunizations})
